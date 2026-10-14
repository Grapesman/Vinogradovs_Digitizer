from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AxisScale, CurveData, DepthCalibration


def build_output_dataframe(
    curves: list[CurveData], depth: DepthCalibration, depth_step: float
) -> pd.DataFrame:
    if depth_step <= 0:
        raise ValueError("Шаг по глубине должен быть > 0")
    usable = [c for c in curves if len(c.points_hv) >= 2]
    if not usable:
        raise ValueError("Нет распознанных кривых для экспорта")

    h_low = min(depth.h1, depth.h2)
    h_high = max(depth.h1, depth.h2)
    n = int(np.floor((h_high - h_low) / depth_step + 1e-9)) + 1
    grid = h_low + np.arange(n + 1, dtype=float) * depth_step
    grid = grid[grid <= h_high + depth_step * 0.25]

    data: dict[str, np.ndarray] = {"H, м": grid}
    for curve in usable:
        arr = np.asarray(curve.points_hv, dtype=float)
        hs = arr[:, 0]
        vals = arr[:, 1]
        order = np.argsort(hs)
        hs, vals = hs[order], vals[order]
        unique_h, unique_idx = np.unique(hs, return_index=True)
        vals = vals[unique_idx]
        col = np.full(grid.shape, np.nan, dtype=float)
        valid = (grid >= unique_h.min()) & (grid <= unique_h.max())
        col[valid] = np.interp(grid[valid], unique_h, vals)
        header = curve.config.name.strip() or "Curve"
        if curve.config.unit.strip():
            header = f"{header}, {curve.config.unit.strip()}"
        base = header
        suffix = 2
        while header in data:
            header = f"{base} ({suffix})"
            suffix += 1
        data[header] = col

    return pd.DataFrame(data)


def _curves_metadata(curves: list[CurveData], scales: list[AxisScale]) -> pd.DataFrame:
    scale_names = {s.id: s.name for s in scales}
    rows = []
    for c in curves:
        cfg = c.config
        q = c.quality or {}
        rows.append({
            "Name": cfg.name,
            "Unit": cfg.unit,
            "Scale": scale_names.get(cfg.scale_id, cfg.scale_id),
            "Scale ID": cfg.scale_id,
            "ROI x": cfg.roi[0],
            "ROI y": cfg.roi[1],
            "ROI width": cfg.roi[2],
            "ROI height": cfg.roi[3],
            "Color samples": len(cfg.effective_colors()),
            "LAB tolerance": cfg.color_tolerance,
            "Precise colour separation": cfg.precise_color,
            "Competitor LAB margin": cfg.competitor_margin,
            "Hue tolerance (OpenCV)": cfg.hue_tolerance,
            "Grid suppression": cfg.remove_grid,
            "Grid strength": cfg.grid_strength,
            "Max jump px": cfg.max_jump_px,
            "Max gap rows": cfg.max_gap_rows,
            "Smooth window": cfg.smooth_window,
            "Detected points": len(c.points_xy),
            "Coverage %": q.get("coverage_pct", np.nan),
            "P95 jump px": q.get("p95_jump_px", np.nan),
            "Max gap rows observed": q.get("max_gap_rows", np.nan),
            "Quality score": q.get("score", np.nan),
        })
    return pd.DataFrame(rows)


def _scales_metadata(scales: list[AxisScale]) -> pd.DataFrame:
    return pd.DataFrame([
        {
            "ID": s.id,
            "Name": s.name,
            "Unit": s.unit,
            "Type": s.scale_type,
            "Min": s.value_min,
            "Max": s.value_max,
            "X1 px": s.x1,
            "X2 px": s.x2,
            "Division": s.division,
            "ROI x": s.roi[0],
            "ROI y": s.roi[1],
            "ROI width": s.roi[2],
            "ROI height": s.roi[3],
        }
        for s in scales
    ])


def export_xlsx(
    path: str | Path,
    df: pd.DataFrame,
    curves: list[CurveData] | None = None,
    scales: list[AxisScale] | None = None,
    depth: DepthCalibration | None = None,
    source_image: str = "",
    depth_step: float | None = None,
) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    curves = curves or []
    scales = scales or []
    project_rows = []
    if source_image:
        project_rows.append(("Source image", source_image))
    if depth is not None:
        project_rows += [
            ("H1, m", depth.h1), ("Y1, px", depth.y1),
            ("H2, m", depth.h2), ("Y2, px", depth.y2),
        ]
    if depth_step is not None:
        project_rows.append(("Depth step, m", depth_step))
    project_rows += [("Scales", len(scales)), ("Curves", len(curves))]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data")
        _curves_metadata(curves, scales).to_excel(writer, index=False, sheet_name="Curves")
        _scales_metadata(scales).to_excel(writer, index=False, sheet_name="Scales")
        pd.DataFrame(project_rows, columns=["Parameter", "Value"]).to_excel(
            writer, index=False, sheet_name="Project"
        )

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="E9EFF7")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="173F7A")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for col_idx, col in enumerate(ws.iter_cols(1, ws.max_column), start=1):
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 36)
    wb.save(path)
